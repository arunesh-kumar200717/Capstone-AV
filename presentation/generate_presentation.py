from __future__ import annotations

from pathlib import Path

from pptx import Presentation

OUTPUT_PATH = Path('presentation/research_paper_answer_bot_presentation.pptx')


def add_title_slide(prs: Presentation, title: str, subtitle: str):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_content_slide(prs: Presentation, title: str, bullets: list[str], notes: str | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.shapes.placeholders[1].text_frame
    for idx, bullet in enumerate(bullets):
        paragraph = body.paragraphs[0] if idx == 0 else body.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


if __name__ == '__main__':
    prs = Presentation()
    add_title_slide(prs, 'Research Paper Answer Bot', 'Analytics Vidhya Capstone | Project Overview')
    add_content_slide(prs, 'Business problem', ['Need for grounded research Q&A over PDFs.', 'Manual reading is slow and error-prone.', 'Researchers need citeable evidence with retrieval support.'], 'Discuss the need for grounded and verifiable answers.')
    add_content_slide(prs, 'Proposed RAG architecture', ['PDF ingestion and chunking', 'Vector storage in ChromaDB', 'Dense and hybrid retrieval', 'Prompted generation with citations'], 'Explain the end-to-end architecture and why each stage matters.')
    add_content_slide(prs, 'Dataset and preprocessing', ['Collect research PDFs', 'Preserve title and page metadata', 'Handle OCR need for scanned PDFs', 'Chunk with overlap for better retrieval'], 'Describe how data is ingested and normalized before retrieval.')
    add_content_slide(prs, 'Chunking methodology', ['Page-level extraction', 'Chunk size and overlap parameters', 'Stable chunk IDs and metadata', 'Indexing at the passage level'], 'Explain design choices and why metadata matters.')
    add_content_slide(prs, 'Embedding model comparison', ['all-MiniLM-L6-v2', 'BAAI/bge-small-en-v1.5', 'Configurable embedding selection', 'Model-specific vector collections'], 'Note that embeddings are compared in the retrieval lab and can be configured per run.')
    add_content_slide(prs, 'Vector DB and retrieval strategies', ['ChromaDB persistent storage', 'Dense retrieval', 'BM25 keyword retrieval', 'Hybrid rank-based fusion'], 'Explain how retrieval quality is improved using multiple strategies.')
    add_content_slide(prs, 'RAG pipeline and prompt design', ['Query embedding and retrieval', 'Context assembly from retrieved passages', 'Grounded answer generation', 'Citation validation'], 'Discuss the importance of limiting answers to evidence in the context.')
    add_content_slide(prs, 'Actual results', ['Project scaffolding and functionality implemented', 'Evaluation scripts and sample query set added', 'Results pending live model-backed benchmarking', 'Use placeholders when experiments are not run'], 'Clearly mark which results were executed and which remain pending.')
    add_content_slide(prs, 'Conversational memory stretch goal', ['Session-based chat history', 'Follow-up question handling', 'Fresh retrieval for each follow-up', 'Source citation cards'], 'Highlight how the assistant maintains context without sacrificing grounding.')
    add_content_slide(prs, 'Challenges and lessons learned', ['Scanned PDFs need OCR', 'Model choice affects retrieval quality', 'Citations must be validated carefully', 'LLM config errors must be surfaced clearly'], 'Discuss limitations and practical constraints.')
    add_content_slide(prs, 'Conclusion and future enhancements', ['Add reranking and live benchmarks', 'Support a broader corpus and stronger evaluation', 'Ship with a notebook or dashboard for final reporting'], 'Emphasize how the system is ready for future extension.')
    prs.save(str(OUTPUT_PATH))
    print(f'PowerPoint presentation generated: {OUTPUT_PATH}')
